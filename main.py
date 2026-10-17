"""Web scraper for books.toscrape.com with a formatted Excel report (USD).

The practice site lists prices in GBP (£). This tool converts them to USD with a
configurable exchange rate and exports a styled .xlsx with a summary sheet.

Basic usage:
    python main.py
    python main.py --max-pages 5 --output books.xlsx
    python main.py --rate 1.34 --include-gbp
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
import time
from datetime import datetime
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

START_URL = "http://books.toscrape.com/catalogue/page-1.html"
TIMEOUT = 10  # seconds
DEFAULT_RATE = float(os.getenv("GBP_USD_RATE", "1.34"))  # GBP -> USD; update before real use
SHEET_NAME = "Books"
SUMMARY_SHEET = "Summary"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
USD_FORMAT = '"$"#,##0.00'  # shows "$" regardless of regional settings

log = logging.getLogger("scraper")


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------
def build_session() -> requests.Session:
    """HTTP session with its own User-Agent and automatic retries."""
    session = requests.Session()
    session.headers.update(
        {"User-Agent": "Mozilla/5.0 (compatible; python-web-scraper-excel/1.1)"}
    )
    retry = Retry(total=3, backoff_factor=1, status_forcelist=(429, 500, 502, 503, 504))
    adapter = HTTPAdapter(max_retries=retry)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session


def parse_price(text: str) -> float:
    """Converts '£51.77' into 51.77."""
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        raise ValueError(f"Could not parse price: {text!r}")
    return float(match.group())


def parse_books(soup: BeautifulSoup, page_url: str) -> list[dict]:
    """Extracts the books from one catalog page (prices stay in GBP here)."""
    books = []
    for article in soup.select("article.product_pod"):
        rating_tag = article.select_one("p.star-rating")
        rating_word = next(
            (c for c in (rating_tag.get("class", []) if rating_tag else []) if c in RATINGS),
            None,
        )
        books.append(
            {
                "Title": article.h3.a["title"],
                "price_gbp": parse_price(article.select_one("p.price_color").text),
                "Rating (1-5)": RATINGS.get(rating_word),
                "Availability": article.select_one("p.instock.availability").text.strip(),
                "URL": urljoin(page_url, article.h3.a["href"]),
            }
        )
    return books


def scrape_books(
    start_url: str = START_URL, max_pages: int | None = None, delay: float = 0.5
) -> list[dict]:
    """Walks the catalog following the 'next' link until the end."""
    session = build_session()
    data: list[dict] = []
    url: str | None = start_url
    pages = 0

    while url and (max_pages is None or pages < max_pages):
        log.info("Reading %s", url)
        try:
            response = session.get(url, timeout=TIMEOUT)
            response.raise_for_status()
        except requests.RequestException as exc:
            log.error("Could not access %s: %s", url, exc)
            break

        # The site does not declare its encoding; without this '£' arrives as 'Â£'.
        response.encoding = "utf-8"
        soup = BeautifulSoup(response.text, "html.parser")

        data.extend(parse_books(soup, response.url))
        pages += 1

        next_link = soup.select_one("li.next > a")
        url = urljoin(response.url, next_link["href"]) if next_link else None
        if url:
            time.sleep(delay)  # be polite to the server

    log.info("Pages read: %d | Books extracted: %d", pages, len(data))
    return data


# --------------------------------------------------------------------------
# Transformation
# --------------------------------------------------------------------------
def to_dataframe(data: list[dict], rate: float, include_gbp: bool = False) -> pd.DataFrame:
    """Builds the report table: USD prices, plus the original GBP price if requested."""
    df = pd.DataFrame(data)
    df["Price ($)"] = (df["price_gbp"] * rate).round(2)
    columns = ["Title", "Price ($)"]
    if include_gbp:
        df["Source Price (GBP)"] = df["price_gbp"]
        columns.append("Source Price (GBP)")
    return df[columns + ["Rating (1-5)", "Availability", "URL"]]


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------
def export_to_styled_excel(
    df: pd.DataFrame, output_filename: str, rate: float, source_url: str
) -> None:
    """Exports the data to a styled .xlsx with a summary sheet."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center")
    link_font = Font(color="0563C1", underline="single")

    with pd.ExcelWriter(output_filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=SHEET_NAME)
        ws = writer.sheets[SHEET_NAME]

        for cell in ws[1]:
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center

        columns = list(df.columns)
        price_col = columns.index("Price ($)")
        rating_col = columns.index("Rating (1-5)")
        url_col = columns.index("URL")
        gbp_col = columns.index("Source Price (GBP)") if "Source Price (GBP)" in columns else None

        for row in ws.iter_rows(min_row=2):
            row[price_col].number_format = USD_FORMAT
            row[rating_col].alignment = center
            if gbp_col is not None:
                row[gbp_col].number_format = '"£"#,##0.00'
            url_cell = row[url_col]
            url_cell.hyperlink = url_cell.value
            url_cell.font = link_font

        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            longest = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(
                max(longest + 3, 12), 60
            )

        # ---- Summary sheet ----
        prices = df["Price ($)"]
        stats = [
            ("Run date", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), None),
            ("Source", source_url, None),
            ("Source currency", "GBP (£)", None),
            ("Report currency", "USD ($)", None),
            ("Exchange rate GBP→USD", rate, "0.0000"),
            ("Books extracted", len(df), "#,##0"),
            ("Average price ($)", round(prices.mean(), 2), USD_FORMAT),
            ("Lowest price ($)", prices.min(), USD_FORMAT),
            ("Highest price ($)", prices.max(), USD_FORMAT),
            ("Average rating (1-5)", round(df["Rating (1-5)"].mean(), 2), "0.00"),
            ("Note", "Prices converted from GBP at the rate shown. Excludes taxes and shipping.", None),
        ]
        pd.DataFrame({"Item": [s[0] for s in stats], "Value": [s[1] for s in stats]}).to_excel(
            writer, index=False, sheet_name=SUMMARY_SHEET
        )
        ws2 = writer.sheets[SUMMARY_SHEET]
        for cell in ws2[1]:
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center
        for i, (_, _, fmt) in enumerate(stats, start=2):
            cell = ws2.cell(row=i, column=2)
            cell.alignment = Alignment(horizontal="left", wrap_text=True)
            if fmt:
                cell.number_format = fmt
        ws2.column_dimensions["A"].width = 28
        ws2.column_dimensions["B"].width = 60

    log.info("Data exported to %s", output_filename)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Book scraper with Excel export (USD).")
    parser.add_argument("--url", default=START_URL, help="URL of the first catalog page.")
    parser.add_argument("--max-pages", type=int, default=None, help="Page limit (default: all pages).")
    parser.add_argument("--delay", type=float, default=0.5, help="Seconds to wait between pages.")
    parser.add_argument("--rate", type=float, default=DEFAULT_RATE,
                        help="GBP->USD exchange rate (default: env GBP_USD_RATE or 1.34).")
    parser.add_argument("--include-gbp", action="store_true",
                        help="Also include the original GBP price column.")
    parser.add_argument("--output", default="books_report.xlsx", help="Excel output file name.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    results = scrape_books(args.url, args.max_pages, args.delay)
    if not results:
        log.error("No data extracted; Excel file not generated.")
        return 1

    df = to_dataframe(results, args.rate, args.include_gbp)
    export_to_styled_excel(df, args.output, args.rate, args.url)
    return 0


if __name__ == "__main__":
    sys.exit(main())
